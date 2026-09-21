"""One-step applicant profile creation from CV and supporting documents."""

from __future__ import annotations

import hashlib
import io
import json
import os
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from PyPDF2 import PdfReader

from phd_agent.applicant_context import ApplicantResearchContextService
from phd_agent.db import connect, transaction, utc_now
from phd_agent.documents import DocumentVault
from phd_agent.materials import MaterialStudio
from phd_agent.model_router import ModelConfig
from phd_agent.profile import ApplicantTruth


@dataclass(frozen=True)
class ProfileUpload:
    name: str
    data: bytes


@dataclass(frozen=True)
class ProfileBuildResult:
    context_id: int
    profile_id: int
    profile_version_id: int
    master_cv_version_id: int
    research_track_version_id: int
    summary_path: Path
    documents_added: int
    facts_in_profile: int
    warnings: tuple[str, ...]


SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}
PROFILE_DOCUMENT_TYPES = {
    "CV", "SOP", "PERSONAL_STATEMENT", "RESEARCH_PROPOSAL", "RESEARCH_STATEMENT",
    "DEGREE_CERTIFICATE", "TRANSCRIPT", "MARKSHEET", "ENGLISH_TEST",
    "STANDARDIZED_TEST", "PATENT_DOCUMENT", "PUBLICATION", "CERTIFICATE", "OTHER",
}
SECTION_BY_CATEGORY = {
    "EDUCATION": "Education",
    "RESEARCH_PROJECT": "Research Experience",
    "INDUSTRY_RESEARCH": "Industry and Research Engineering",
    "PATENT": "Patents",
    "PUBLICATION": "Publications",
    "TEACHING": "Teaching",
    "SKILL_METHOD": "Technical Skills",
    "OTHER": "Projects",
}
SUMMARY_HEADINGS = {
    "EDUCATION": "Education",
    "RESEARCH_PROJECT": "Research and projects",
    "INDUSTRY_RESEARCH": "Professional and engineering experience",
    "PATENT": "Patents",
    "PUBLICATION": "Publications",
    "TEACHING": "Teaching",
    "SKILL_METHOD": "Skills and methods",
    "OTHER": "Additional experience",
}


def infer_document_type(filename: str, text: str = "") -> str:
    value = (Path(filename).stem + " " + text[:800]).casefold()
    rules = (
        (("curriculum vitae", "resume", " cv "), "CV"),
        (("transcript", "academic record", "hear"), "TRANSCRIPT"),
        (("marksheet", "mark sheet"), "MARKSHEET"),
        (("degree certificate", "award certificate", "diploma"), "DEGREE_CERTIFICATE"),
        (("research proposal",), "RESEARCH_PROPOSAL"),
        (("statement of purpose", " sop "), "SOP"),
        (("personal statement",), "PERSONAL_STATEMENT"),
        (("research statement",), "RESEARCH_STATEMENT"),
        (("ielts", "toefl", "english test"), "ENGLISH_TEST"),
        (("patent",), "PATENT_DOCUMENT"),
        (("publication", "paper", "manuscript"), "PUBLICATION"),
        (("certificate",), "CERTIFICATE"),
    )
    padded = f" {value} "
    for signals, document_type in rules:
        if any(signal in padded for signal in signals):
            return document_type
    return "OTHER"


def extract_text(data: bytes, filename: str) -> tuple[str, int, str]:
    suffix = Path(filename).suffix.casefold()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"{filename}: use PDF, DOCX, TXT, or Markdown")
    if suffix == ".pdf":
        pages = [page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages]
        text = "\n\f\n".join(pages)
        method = "PyPDF2-all-pages"
    elif suffix == ".docx":
        document = DocxDocument(io.BytesIO(data))
        pieces = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                pieces.append(" | ".join(cell.text for cell in row.cells))
        pages = ["\n".join(pieces)]
        text = pages[0]
        method = "python-docx-paragraphs-and-tables"
    else:
        text = data.decode("utf-8", errors="replace")
        pages = [text]
        method = "utf-8-text"
    text = text.replace("\x00", "").strip()
    if len(text) < 20:
        raise ValueError(f"{filename}: no useful text could be extracted")
    return text, len(pages), method


def _normalized(value: str) -> str:
    return re.sub(r"\W+", " ", value.casefold()).strip()


def _category(line: str) -> str:
    value = line.casefold()
    if re.search(r"\b(phd|msc|m\.sc|btech|b\.tech|bachelor|master|degree|university|college|gpa|cgpa)\b", value):
        return "EDUCATION"
    if re.search(r"\b(patent|invention)\b", value):
        return "PATENT"
    if re.search(r"\b(publication|published|journal|conference|doi|manuscript|paper)\b", value):
        return "PUBLICATION"
    if re.search(r"\b(teach|teaching|lecturer|tutor|mentor)\b", value):
        return "TEACHING"
    if re.search(r"\b(python|java|pytorch|tensorflow|sql|llm|rag|machine learning|computer vision|robotics|skill|method)\b", value):
        return "SKILL_METHOD"
    if re.search(r"\b(engineer|employment|company|industry|intern|professional experience)\b", value):
        return "INDUSTRY_RESEARCH"
    if re.search(r"\b(research|project|thesis|dissertation|developed|built|evaluated|implemented|experiment)\b", value):
        return "RESEARCH_PROJECT"
    return "OTHER"


def _deterministic_candidates(text: str, *, limit: int = 36) -> list[dict]:
    candidates = []
    seen = set()
    for raw in text.splitlines():
        line = re.sub(r"\s+", " ", raw).strip(" •\t-|:")
        if not 24 <= len(line) <= 700 or re.fullmatch(r"(?:page\s*)?\d+", line, re.I):
            continue
        key = _normalized(line)
        if len(key.split()) < 4 or key in seen:
            continue
        seen.add(key)
        future = bool(re.search(r"\b(aim|hope|intend|plan|would like|future research|interested in)\b", line, re.I))
        statement = line
        if future and not re.search(r"\b(i aim|i hope|i intend|i plan|i would like|my goal|future research|i am interested in)\b", line, re.I):
            statement = "I aim to " + line[0].lower() + line[1:]
        candidates.append({
            "statement": statement,
            "category": _category(line),
            "classification": "ASPIRATION" if future else "FACT",
            "normalized_claim_type": "document_statement",
            "source_quote": line,
        })
        if len(candidates) >= limit:
            break
    if not candidates:
        excerpt = re.sub(r"\s+", " ", text)[:600].strip()
        candidates.append({
            "statement": excerpt, "category": "OTHER", "classification": "FACT",
            "normalized_claim_type": "document_statement", "source_quote": excerpt,
        })
    return candidates


class ProfileWorkspace:
    """Turns uploaded source documents into a usable context with one operator action."""

    reviewer = "Profile workspace"

    def __init__(self, db_path: Path | str):
        self.db_path = Path(db_path)
        self.vault = DocumentVault(self.db_path)
        self.truth = ApplicantTruth(self.db_path, self.vault)
        self.studio = MaterialStudio(self.db_path, self.vault)
        self.contexts = ApplicantResearchContextService(self.db_path)
        self.summary_dir = self.db_path.parent / "profile_summaries"
        self.summary_dir.mkdir(parents=True, exist_ok=True)

    def source_documents(self) -> list[dict]:
        with connect(self.db_path) as db:
            return [dict(row) for row in db.execute("""SELECT d.*,v.id AS version_id,v.original_filename,
                v.approval_state,v.verification_state,v.storage_key,v.sha256
                FROM documents d JOIN document_versions v ON v.document_id=d.id
                WHERE d.document_class='SOURCE' AND d.document_type IN (%s)
                  AND v.version_number=(SELECT MAX(v2.version_number) FROM document_versions v2 WHERE v2.document_id=d.id)
                ORDER BY d.id""" % ",".join("?" for _ in PROFILE_DOCUMENT_TYPES), tuple(sorted(PROFILE_DOCUMENT_TYPES)))]

    def latest_summary(self) -> tuple[Path, str] | None:
        paths = sorted(self.summary_dir.glob("applicant-profile-v*.md"),
                       key=lambda path: path.stat().st_mtime, reverse=True)
        if not paths:
            return None
        return paths[0], paths[0].read_text(encoding="utf-8")

    def suggested_name(self) -> str:
        with connect(self.db_path) as db:
            row = db.execute("SELECT owner_name FROM applicant_profiles ORDER BY id DESC LIMIT 1").fetchone()
        return row["owner_name"] if row else (os.getenv("USER_NAME") or "").strip()

    def suggested_focus(self) -> str:
        latest = self.contexts.latest()
        if latest:
            return latest["context"]["research_track"]["title"]
        return ""

    def _profile_id(self, owner_name: str) -> int:
        with connect(self.db_path) as db:
            row = db.execute("SELECT id FROM applicant_profiles WHERE lower(owner_name)=lower(?) ORDER BY id DESC LIMIT 1",
                             (owner_name,)).fetchone()
        return row["id"] if row else self.truth.create_profile(owner_name)

    def _record_extraction(self, version_id: int, text: str, page_count: int, method: str) -> int:
        with connect(self.db_path) as db:
            existing = db.execute("SELECT id FROM applicant_document_extractions WHERE document_version_id=?",
                                  (version_id,)).fetchone()
        if existing:
            return existing["id"]
        with transaction(self.db_path) as db:
            return db.execute("""INSERT INTO applicant_document_extractions
                (document_version_id,extraction_method,extracted_at,page_count,extracted_text,text_sha256,notes)
                VALUES(?,?,?,?,?,?,?)""", (
                version_id, method, utc_now(), page_count, text,
                hashlib.sha256(text.encode()).hexdigest(), "Created by the simple profile workspace",
            )).lastrowid

    def _approve_document(self, version_id: int) -> None:
        version = self.vault.get_version(version_id)
        if version["approval_state"] == "APPROVED":
            return
        if version["document_type"] in {"TRANSCRIPT", "DEGREE_CERTIFICATE"}:
            self.vault.set_verification(version_id, "VERIFIED", self.reviewer)
        self.vault.set_approval(version_id, True, self.reviewer)

    def _new_claims(self, profile_id: int, version_id: int, extraction_id: int, text: str,
                    *, api_key: str = "") -> tuple[list[int], list[str]]:
        with connect(self.db_path) as db:
            existing = [row["id"] for row in db.execute("""SELECT cr.id FROM claim_revisions cr
                JOIN claims c ON c.id=cr.claim_id
                WHERE cr.extraction_id=? AND cr.review_status='APPROVED' AND c.profile_id=?""",
                (extraction_id, profile_id))]
        if existing:
            return existing, []
        warnings = []
        if api_key:
            try:
                model = ModelConfig.from_environ().luna
                self.truth.extract_candidates(profile_id, extraction_id, api_key=api_key, model=model)
            except Exception as error:
                warnings.append(f"Structured extraction was unavailable; used local parsing ({type(error).__name__}).")
        with connect(self.db_path) as db:
            pending = [dict(row) for row in db.execute("""SELECT cr.* FROM claim_revisions cr
                JOIN claims c ON c.id=cr.claim_id
                WHERE cr.extraction_id=? AND cr.review_status='PENDING' AND c.profile_id=?""",
                (extraction_id, profile_id))]
        if not pending:
            for item in _deterministic_candidates(text):
                revision_id = self.truth.create_claim(
                    profile_id, item["category"], item["statement"], item["normalized_claim_type"],
                    item["classification"], source_document_version_id=version_id,
                    source_location="exact extracted line", extraction_id=extraction_id,
                    confidence=1.0, notes="Exact source quote: " + item["source_quote"],
                )
                pending.append(self.truth._claim_revision(revision_id))
        approved = []
        for claim in pending:
            self.truth.review_claim(
                claim["id"], True, self.reviewer, verification_state="VERIFIED",
                for_application=True, for_outreach=True,
                notes=(claim["notes"] or "") + "\nIncluded through the simple profile workspace.",
            )
            approved.append(claim["id"])
        return approved, warnings

    def _profile_version(self, profile_id: int, claim_ids: list[int]) -> int:
        claim_ids = sorted(set(claim_ids))
        with connect(self.db_path) as db:
            latest = db.execute("""SELECT id FROM profile_versions WHERE profile_id=? AND approval_state='APPROVED'
                ORDER BY version_number DESC LIMIT 1""", (profile_id,)).fetchone()
            if latest:
                current = [row[0] for row in db.execute(
                    "SELECT claim_revision_id FROM profile_version_claims WHERE profile_version_id=? ORDER BY claim_revision_id",
                    (latest["id"],))]
                if current == claim_ids:
                    return latest["id"]
        return self.truth.create_profile_version(
            profile_id, claim_ids, approve=True, reviewer=self.reviewer,
            notes="Built from uploaded applicant documents in the simple workspace",
        )

    def _track(self, profile_id: int, research_focus: str, supporting_ids: list[int]) -> int:
        focus = research_focus.strip()
        with connect(self.db_path) as db:
            existing = db.execute("""SELECT v.id,t.title FROM research_tracks t
                JOIN research_track_versions v ON v.track_id=t.id
                WHERE t.profile_id=? AND t.status='ACTIVE' AND v.approval_state='APPROVED'
                ORDER BY v.id DESC LIMIT 1""", (profile_id,)).fetchone()
        if existing and (not focus or _normalized(existing["title"]) == _normalized(focus)):
            return existing["id"]
        focus = focus or "Research aligned with the applicant's demonstrated experience"
        _, version_id = self.truth.create_track(
            profile_id, focus,
            research_problem=focus,
            research_questions=f"How can {focus.rstrip('.').lower()} be advanced and evaluated rigorously?",
            proposed_methodology="Use methods suited to each target programme and grounded in the applicant's demonstrated experience.",
            evaluation_strategy="Define measurable outcomes, baselines, failure cases, and reproducible evaluation for each proposed project.",
            expected_contribution=f"A rigorous contribution to {focus.rstrip('.').lower()}.",
            supporting_claim_revision_ids=supporting_ids[:8],
        )
        self.truth.approve_track(version_id, self.reviewer)
        return version_id

    def _master_cv(self, profile_id: int, profile_version_id: int, claims: list[dict]) -> int:
        with connect(self.db_path) as db:
            existing = db.execute("""SELECT id FROM master_cv_versions WHERE profile_id=? AND profile_version_id=?
                AND approval_state='APPROVED' ORDER BY version_number DESC LIMIT 1""",
                (profile_id, profile_version_id)).fetchone()
        if existing:
            return existing["id"]
        grouped: dict[str, list[dict]] = {}
        for claim in claims:
            section = SECTION_BY_CATEGORY.get(claim["category"], "Projects")
            grouped.setdefault(section, []).append({
                "text": claim["claim_text"], "claim_revision_ids": [claim["id"]],
            })
        sections = [{"name": section, "bullets": bullets[:12]} for section, bullets in grouped.items() if bullets]
        master_id = self.studio.create_master_cv(profile_version_id, sections)
        self.studio.review_master_cv(master_id, self.reviewer, True)
        return master_id

    def _summary(self, owner_name: str, profile_version_id: int, research_focus: str,
                 claims: list[dict], documents: list[dict]) -> tuple[Path, str]:
        lines = [f"# Applicant profile — {owner_name}", "", f"_Profile version {profile_version_id}_", "",
                 "## Research direction", "", research_focus.strip() or "To be refined for each application.", ""]
        by_category: dict[str, list[dict]] = {}
        for claim in claims:
            by_category.setdefault(claim["category"], []).append(claim)
        for category, heading in SUMMARY_HEADINGS.items():
            items = by_category.get(category, [])
            if not items:
                continue
            lines.extend([f"## {heading}", ""])
            for claim in items[:16]:
                source = claim.get("source_filename") or "uploaded document"
                lines.append(f"- {claim['claim_text']}  ")
                lines.append(f"  _Source: {source}_")
            lines.append("")
        lines.extend(["## Source documents", ""])
        for document in documents:
            lines.append(f"- **{document['original_filename']}** — {document['document_type'].replace('_', ' ').title()}")
        lines.extend(["", "## How this file is used", "",
                      "This summary is a readable index of source-backed applicant information. "
                      "Application documents and professor outreach still retrieve the relevant underlying claims and source files.", ""])
        markdown = "\n".join(lines)
        path = self.summary_dir / f"applicant-profile-v{profile_version_id}.md"
        path.write_text(markdown, encoding="utf-8")
        summary_upload = self.vault.upload(
            markdown.encode(), path.name, "GENERATED", "OTHER", "Applicant profile summary",
            notes=f"Readable summary for profile version {profile_version_id}",
        )
        self._approve_document(summary_upload["version_id"],)
        return path, markdown

    def build(self, owner_name: str, research_focus: str, uploads: list[ProfileUpload] | None = None,
              *, api_key: str = "") -> ProfileBuildResult:
        owner = owner_name.strip()
        if not owner:
            raise ValueError("Add your name so the profile can be created")
        profile_id = self._profile_id(owner)
        added = 0
        warnings: list[str] = []
        for upload in uploads or []:
            if not upload.data:
                continue
            text, _, _ = extract_text(upload.data, upload.name)
            document_type = infer_document_type(upload.name, text)
            saved = self.vault.upload(
                upload.data, upload.name, "SOURCE", document_type,
                Path(upload.name).stem.replace("_", " ").replace("-", " ").strip().title(),
                sensitivity="CONFIDENTIAL" if document_type in {"TRANSCRIPT", "MARKSHEET", "DEGREE_CERTIFICATE"} else "NORMAL",
                notes="Added through the simple profile workspace",
            )
            self._approve_document(saved["version_id"])
            added += int(saved["status"] == "created")
        documents = self.source_documents()
        if not documents:
            raise ValueError("Add at least one CV or supporting document")
        for document in documents:
            self._approve_document(document["version_id"])
            try:
                data = self.vault.storage.get(document["storage_key"])
                text, page_count, method = extract_text(data, document["original_filename"])
            except ValueError as error:
                warnings.append(str(error))
                continue
            extraction_id = self._record_extraction(document["version_id"], text, page_count, method)
            _, extraction_warnings = self._new_claims(
                profile_id, document["version_id"], extraction_id, text, api_key=api_key)
            warnings.extend(extraction_warnings)
        claims = [claim for claim in self.truth.list_claims(profile_id)
                  if claim["review_status"] == "APPROVED" and claim["approved_for_application"]]
        if not claims:
            raise ValueError("The uploaded documents did not contain enough readable profile information")
        claim_ids = [claim["id"] for claim in claims]
        profile_version_id = self._profile_version(profile_id, claim_ids)
        supporting = [claim["id"] for claim in claims
                      if claim["classification"] == "FACT" and claim["category"] in {
                          "RESEARCH_PROJECT", "INDUSTRY_RESEARCH", "SKILL_METHOD", "PUBLICATION", "PATENT"}]
        supporting = supporting or [claim["id"] for claim in claims if claim["classification"] == "FACT"]
        track_version_id = self._track(profile_id, research_focus, supporting)
        master_cv_version_id = self._master_cv(profile_id, profile_version_id, claims)
        context = self.contexts.build(profile_version_id, master_cv_version_id, track_version_id)
        focus = context["context"]["research_track"]["title"]
        summary_path, _ = self._summary(owner, profile_version_id, focus, claims, documents)
        return ProfileBuildResult(
            context["id"], profile_id, profile_version_id, master_cv_version_id,
            track_version_id, summary_path, added, len(claims), tuple(dict.fromkeys(warnings)),
        )
