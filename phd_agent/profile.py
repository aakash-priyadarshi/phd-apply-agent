"""Versioned applicant truth, candidate extraction, and research directions."""

from __future__ import annotations

import hashlib
import io
import json
import re
from pathlib import Path
from typing import Literal

from docx import Document as DocxDocument
from openai import OpenAI
from pydantic import BaseModel, Field
from PyPDF2 import PdfReader

from phd_agent.db import connect, migrate, transaction, utc_now
from phd_agent.documents import DocumentVault


CATEGORIES = (
    "EDUCATION", "RESEARCH_PROJECT", "INDUSTRY_RESEARCH", "PATENT",
    "PUBLICATION", "TEACHING", "SKILL_METHOD", "OTHER",
)
CLASSIFICATIONS = ("FACT", "INFERENCE", "ASPIRATION")
INGEST_TYPES = {
    "CV", "SOP", "PERSONAL_STATEMENT", "RESEARCH_PROPOSAL", "RESEARCH_STATEMENT",
    "PATENT_DOCUMENT", "PUBLICATION", "CERTIFICATE",
}
TRACK_FIELDS = (
    "research_problem", "motivation", "research_gap", "research_questions",
    "hypotheses", "proposed_methodology", "evaluation_strategy",
    "possible_datasets", "expected_contribution", "related_projects",
    "prior_work", "limitations", "open_questions", "notes",
)
EXTRACTION_PROMPT_VERSION = "candidate-claims-v1"


class ExtractedField(BaseModel):
    key: str
    value: str


class CandidateStatement(BaseModel):
    statement: str
    category: Literal[
        "EDUCATION", "RESEARCH_PROJECT", "INDUSTRY_RESEARCH", "PATENT",
        "PUBLICATION", "TEACHING", "SKILL_METHOD", "OTHER",
    ]
    normalized_claim_type: str
    classification: Literal["FACT", "INFERENCE", "ASPIRATION"]
    source_quote: str
    source_location: str
    confidence: float = Field(ge=0, le=1)
    fields: list[ExtractedField]


class CandidateBatch(BaseModel):
    candidates: list[CandidateStatement]


def _json(value) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _text_normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def _overlapping_chunks(text: str, size: int = 10000, overlap: int = 1000) -> list[str]:
    if not text:
        return []
    if len(text) <= size:
        return [text]
    step = max(1, size - overlap)
    chunks = []
    start = 0
    while start < len(text):
        chunks.append(text[start:start + size])
        if start + size >= len(text):
            break
        start += step
    return chunks


class ApplicantTruth:
    def __init__(self, db_path: Path | str, vault: DocumentVault | None = None):
        self.db_path = Path(db_path)
        migrate(self.db_path)
        self.vault = vault or DocumentVault(self.db_path)

    def create_profile(self, owner_name: str) -> int:
        if not owner_name.strip():
            raise ValueError("Profile owner is required")
        now = utc_now()
        with transaction(self.db_path) as db:
            return db.execute(
                "INSERT INTO applicant_profiles(owner_name,created_at,updated_at) VALUES(?,?,?)",
                (owner_name.strip(), now, now),
            ).lastrowid

    def list_profiles(self) -> list[dict]:
        with connect(self.db_path) as db:
            return [dict(r) for r in db.execute("SELECT * FROM applicant_profiles ORDER BY id")]

    def _claim_revision(self, revision_id: int) -> dict:
        with connect(self.db_path) as db:
            row = db.execute("""SELECT cr.*, c.profile_id, c.category FROM claim_revisions cr
                JOIN claims c ON c.id = cr.claim_id WHERE cr.id = ?""", (revision_id,)).fetchone()
        if not row:
            raise ValueError("Claim revision does not exist")
        return dict(row)

    def create_claim(
        self, profile_id: int, category: str, claim_text: str,
        normalized_claim_type: str, classification: str,
        *, structured_data: dict | None = None, source_document_version_id: int | None = None,
        source_evidence_id: int | None = None, source_location: str | None = None,
        extraction_id: int | None = None, confidence: float | None = None,
        notes: str = "",
    ) -> int:
        if category not in CATEGORIES or classification not in CLASSIFICATIONS:
            raise ValueError("Invalid claim category or classification")
        if not claim_text.strip() or not normalized_claim_type.strip():
            raise ValueError("Claim text and normalized type are required")
        now = utc_now()
        with transaction(self.db_path) as db:
            claim_id = db.execute(
                "INSERT INTO claims(profile_id,category,created_at) VALUES(?,?,?)",
                (profile_id, category, now),
            ).lastrowid
            return db.execute("""INSERT INTO claim_revisions
                (claim_id,version_number,claim_text,normalized_claim_type,structured_data_json,
                 source_document_version_id,source_evidence_id,source_location,extraction_id,
                 classification,confidence,created_at,notes)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""", (
                claim_id, 1, claim_text.strip(), normalized_claim_type.strip(),
                _json(structured_data or {}), source_document_version_id,
                source_evidence_id, source_location, extraction_id, classification,
                confidence, now, notes,
            )).lastrowid

    def revise_claim(self, claim_id: int, **changes) -> int:
        allowed = {
            "claim_text", "normalized_claim_type", "structured_data_json",
            "source_document_version_id", "source_evidence_id", "source_location",
            "classification", "confidence", "verification_state", "notes",
        }
        if set(changes) - allowed:
            raise ValueError("Unsupported claim revision field")
        with transaction(self.db_path) as db:
            latest = db.execute("""SELECT * FROM claim_revisions WHERE claim_id = ?
                ORDER BY version_number DESC LIMIT 1""", (claim_id,)).fetchone()
            if not latest:
                raise ValueError("Claim does not exist")
            values = {field: latest[field] for field in allowed}
            values.update(changes)
            if values["classification"] not in CLASSIFICATIONS:
                raise ValueError("Invalid classification")
            if not values["claim_text"].strip():
                raise ValueError("Claim text is required")
            return db.execute("""INSERT INTO claim_revisions
                (claim_id,version_number,claim_text,normalized_claim_type,structured_data_json,
                 source_document_version_id,source_evidence_id,source_location,classification,
                 confidence,verification_state,created_at,notes)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""", (
                claim_id, latest["version_number"] + 1, values["claim_text"].strip(),
                values["normalized_claim_type"], values["structured_data_json"],
                values["source_document_version_id"], values["source_evidence_id"],
                values["source_location"], values["classification"], values["confidence"],
                values["verification_state"], utc_now(), values["notes"],
            )).lastrowid

    def review_claim(
        self, revision_id: int, approve: bool, reviewer: str,
        *, for_application: bool = False, for_outreach: bool = False,
        verification_state: str = "UNVERIFIED", notes: str | None = None,
    ) -> None:
        revision = self._claim_revision(revision_id)
        if revision["review_status"] != "PENDING":
            raise ValueError("Reviewed claim revisions cannot be changed; create a new revision")
        if verification_state not in {"UNVERIFIED", "VERIFIED", "CONFLICT"}:
            raise ValueError("Invalid verification state")
        has_source = bool(revision["source_document_version_id"] or revision["source_evidence_id"])
        if approve:
            if not reviewer.strip():
                raise ValueError("Reviewer is required")
            if revision["classification"] in {"FACT", "INFERENCE"} and not has_source:
                raise ValueError("FACT and INFERENCE claims need linked evidence before approval")
            if revision["source_document_version_id"]:
                version = self.vault.get_version(revision["source_document_version_id"])
                if not version or version["approval_state"] != "APPROVED":
                    raise ValueError("The source document version must be approved before claim approval")
                if not self.vault.storage.verify_hash(version["storage_key"], version["sha256"]):
                    raise OSError("Source document file is missing or has changed")
            if revision["source_evidence_id"] and revision["classification"] == "FACT":
                with connect(self.db_path) as db:
                    evidence = db.execute("SELECT verification_state FROM source_evidence WHERE id=?",
                                          (revision["source_evidence_id"],)).fetchone()
                if not evidence or evidence["verification_state"] != "VERIFIED":
                    raise ValueError("Factual source evidence must be reviewed before claim approval")
            if verification_state != "VERIFIED" and revision["classification"] == "FACT":
                raise ValueError("A factual claim must be verified before approval")
            if revision["classification"] == "ASPIRATION" and not self._aspiration_wording(revision["claim_text"]):
                raise ValueError("An aspiration must be worded as a future aim, not achieved experience")
        with transaction(self.db_path) as db:
            db.execute("""UPDATE claim_revisions SET review_status = ?, verification_state = ?,
                approved_for_application = ?, approved_for_outreach = ?, approved_at = ?,
                approved_by = ?, notes = ? WHERE id = ?""", (
                "APPROVED" if approve else "REJECTED", verification_state,
                int(approve and for_application), int(approve and for_outreach),
                utc_now() if approve else None, reviewer.strip() if approve else None,
                notes if notes is not None else revision["notes"], revision_id,
            ))

    @staticmethod
    def _aspiration_wording(text: str) -> bool:
        lowered = text.casefold()
        if re.search(
            r"\bmy (?:aim|goal|objective|ambition|aspiration|research interests?)\s+"
            r"(?:was|were|had been)\b",
            lowered,
        ):
            return False
        return any(re.search(rf"\b{re.escape(phrase)}\b", lowered) for phrase in (
            "i aim", "i hope", "i want", "i plan", "i intend", "i propose",
            "i would like", "i seek", "i aspire", "i wish", "i am eager",
            "i am keen", "i look forward", "i envision", "i expect to",
            "i am motivated to", "i am determined to", "i am committed to",
            "this motivates me to", "future research", "my aim", "my goal",
            "my objective", "my ambition", "my aspiration", "my research interest",
            "my research interests", "i am interested in",
        ))

    def list_claims(self, profile_id: int, latest_only: bool = True) -> list[dict]:
        where = "AND cr.version_number = (SELECT MAX(x.version_number) FROM claim_revisions x WHERE x.claim_id = c.id)" if latest_only else ""
        with connect(self.db_path) as db:
            return [dict(r) for r in db.execute(f"""SELECT cr.*, c.profile_id, c.category,
                dv.original_filename AS source_filename, se.canonical_url AS source_url
                FROM claims c JOIN claim_revisions cr ON cr.claim_id = c.id
                LEFT JOIN document_versions dv ON dv.id = cr.source_document_version_id
                LEFT JOIN source_evidence se ON se.id = cr.source_evidence_id
                WHERE c.profile_id = ? {where} ORDER BY c.id, cr.version_number""", (profile_id,))]

    def create_profile_version(
        self, profile_id: int, claim_revision_ids: list[int],
        *, approve: bool = False, reviewer: str | None = None, notes: str = "",
    ) -> int:
        if not claim_revision_ids:
            raise ValueError("Select at least one claim")
        snapshots = []
        for revision_id in sorted(set(claim_revision_ids)):
            revision = self._claim_revision(revision_id)
            if revision["profile_id"] != profile_id:
                raise ValueError("Claim belongs to another profile")
            if approve and revision["review_status"] != "APPROVED":
                raise ValueError("Approved profile snapshots may include only approved claims")
            snapshots.append(revision)
        if approve and not (reviewer or "").strip():
            raise ValueError("Reviewer is required")
        encoded = [_json(snapshot) for snapshot in snapshots]
        snapshot_hash = hashlib.sha256(_json(encoded).encode("utf-8")).hexdigest()
        source_ids = sorted({r["source_document_version_id"] for r in snapshots if r["source_document_version_id"]})
        now = utc_now()
        with transaction(self.db_path) as db:
            version = db.execute(
                "SELECT COALESCE(MAX(version_number),0)+1 FROM profile_versions WHERE profile_id = ?",
                (profile_id,),
            ).fetchone()[0]
            profile_version_id = db.execute("""INSERT INTO profile_versions
                (profile_id,version_number,approval_state,source_document_version_ids_json,
                 claim_snapshot_sha256,notes,created_at,approved_at,approved_by)
                VALUES(?,?,?,?,?,?,?,?,?)""", (
                profile_id, version, "DRAFT", _json(source_ids), snapshot_hash,
                notes, now, None, None,
            )).lastrowid
            db.executemany("""INSERT INTO profile_version_claims
                (profile_version_id,claim_revision_id,claim_snapshot_json) VALUES(?,?,?)""", [
                (profile_version_id, snapshot["id"], encoded[i]) for i, snapshot in enumerate(snapshots)
            ])
            if approve:
                db.execute("""UPDATE profile_versions SET approval_state='APPROVED',
                    approved_at=?, approved_by=? WHERE id=?""", (now, reviewer.strip(), profile_version_id))
            return profile_version_id

    def list_profile_versions(self, profile_id: int) -> list[dict]:
        with connect(self.db_path) as db:
            return [dict(r) for r in db.execute(
                "SELECT * FROM profile_versions WHERE profile_id=? ORDER BY version_number DESC", (profile_id,)
            )]

    def profile_snapshot(self, version_id: int) -> list[dict]:
        with connect(self.db_path) as db:
            return [json.loads(r[0]) for r in db.execute(
                "SELECT claim_snapshot_json FROM profile_version_claims WHERE profile_version_id=? ORDER BY claim_revision_id",
                (version_id,),
            )]

    def ingest_document(self, version_id: int) -> tuple[int, list[str]]:
        version = self.vault.get_version(version_id)
        if not version or version["document_type"] not in INGEST_TYPES:
            raise ValueError("This document type is not eligible for applicant research ingestion")
        if version["approval_state"] != "APPROVED":
            raise ValueError("Approve the source document version before ingestion")
        if not self.vault.storage.verify_hash(version["storage_key"], version["sha256"]):
            raise OSError("Document file is missing or has changed")
        data = self.vault.storage.get(version["storage_key"])
        filename = version["original_filename"].lower()
        if filename.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(data))
            pages = [page.extract_text() or "" for page in reader.pages]
            method = "PyPDF2-all-pages"
        elif filename.endswith(".docx"):
            doc = DocxDocument(io.BytesIO(data))
            pieces = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    pieces.append(" | ".join(cell.text for cell in row.cells))
            pages = ["\n".join(pieces)]
            method = "python-docx-paragraphs-and-tables"
        else:
            raise ValueError("Only PDF and DOCX ingestion is supported")
        text = "\n\f\n".join(pages)
        if not text.strip():
            raise ValueError("The document contains no extractable text")
        with transaction(self.db_path) as db:
            extraction_id = db.execute("""INSERT INTO applicant_document_extractions
                (document_version_id,extraction_method,extracted_at,page_count,
                 extracted_text,text_sha256)
                VALUES(?,?,?,?,?,?)""", (
                version_id, method, utc_now(), len(pages), text,
                hashlib.sha256(text.encode("utf-8")).hexdigest(),
            )).lastrowid
        return extraction_id, pages

    def extract_candidates(
        self, profile_id: int, extraction_id: int, *, api_key: str,
        model: str = "gpt-4o-mini", client=None,
    ) -> dict:
        if not api_key and client is None:
            raise ValueError("An OpenAI API key is required for candidate extraction")
        with connect(self.db_path) as db:
            extraction = db.execute(
                "SELECT * FROM applicant_document_extractions WHERE id = ?", (extraction_id,)
            ).fetchone()
        if not extraction:
            raise ValueError("Extraction does not exist")
        pages = extraction["extracted_text"].split("\n\f\n")
        api = client or OpenAI(api_key=api_key)
        created = 0
        rejected = 0
        for index, page in enumerate(pages, 1):
            if not page.strip():
                continue
            # Every page is processed; long pages are split without dropping a section.
            chunks = _overlapping_chunks(page)
            seen_statements: set[str] = set()
            for chunk_index, chunk in enumerate(chunks, 1):
                response = api.responses.parse(
                    model=model,
                    instructions=(
                        "Extract candidate applicant statements only from the supplied text. "
                        "Return an exact source_quote copied from the text for every statement. "
                        "Distinguish completed facts, inferences, and future aspirations. "
                        "Never turn a patent application into a grant, a manuscript into a publication, "
                        "or tool use into research expertise. Do not approve any claim."
                    ),
                    input=f"Document page {index}, chunk {chunk_index}:\n{chunk}",
                    text_format=CandidateBatch,
                    store=False,
                )
                batch = response.output_parsed
                if not isinstance(batch, CandidateBatch):
                    raise ValueError("Structured candidate extraction returned no parsed result")
                for candidate in batch.candidates:
                    statement_key = _text_normalized(candidate.statement)
                    if statement_key in seen_statements:
                        continue
                    seen_statements.add(statement_key)
                    supported = bool(candidate.source_quote.strip()) and (
                        _text_normalized(candidate.source_quote) in _text_normalized(chunk)
                    )
                    revision_id = self.create_claim(
                        profile_id, candidate.category, candidate.statement,
                        candidate.normalized_claim_type, candidate.classification,
                        structured_data={field.key: field.value for field in candidate.fields},
                        source_document_version_id=extraction["document_version_id"],
                        source_location=candidate.source_location or f"page {index}, chunk {chunk_index}",
                        extraction_id=extraction_id, confidence=candidate.confidence,
                        notes="Exact source quote: " + candidate.source_quote,
                    )
                    if supported:
                        created += 1
                    else:
                        self.review_claim(
                            revision_id, False, "Extractor", notes="Rejected: source quote not found in document chunk"
                        )
                        rejected += 1
        with transaction(self.db_path) as db:
            db.execute("""UPDATE applicant_document_extractions
                SET model_used=?,prompt_version=?,candidate_count=? WHERE id=?""", (
                model, EXTRACTION_PROMPT_VERSION, created + rejected, extraction_id,
            ))
        return {"pending": created, "rejected_unsupported": rejected}

    def create_track(self, profile_id: int, title: str, priority: int = 3, **fields) -> tuple[int, int]:
        if not title.strip() or not 1 <= priority <= 5:
            raise ValueError("Track title and priority 1–5 are required")
        if set(fields) - {*TRACK_FIELDS, "supporting_claim_revision_ids"}:
            raise ValueError("Unsupported research-track field")
        now = utc_now()
        with transaction(self.db_path) as db:
            track_id = db.execute("""INSERT INTO research_tracks
                (profile_id,title,priority,created_at,updated_at) VALUES(?,?,?,?,?)""",
                (profile_id, title.strip(), priority, now, now),
            ).lastrowid
        version_id = self.revise_track(track_id, **fields)
        return track_id, version_id

    def revise_track(self, track_id: int, **changes) -> int:
        if set(changes) - {*TRACK_FIELDS, "supporting_claim_revision_ids"}:
            raise ValueError("Unsupported research-track field")
        with transaction(self.db_path) as db:
            track = db.execute("SELECT * FROM research_tracks WHERE id=?", (track_id,)).fetchone()
            if not track:
                raise ValueError("Research track does not exist")
            previous = db.execute("""SELECT * FROM research_track_versions WHERE track_id=?
                ORDER BY version_number DESC LIMIT 1""", (track_id,)).fetchone()
            values = {field: (previous[field] if previous else "") for field in TRACK_FIELDS}
            values.update({k: v for k, v in changes.items() if k in TRACK_FIELDS})
            claim_ids = changes.get("supporting_claim_revision_ids")
            if claim_ids is None:
                claim_ids = json.loads(previous["supporting_claim_revision_ids_json"]) if previous else []
            for revision_id in claim_ids:
                row = db.execute("""SELECT c.profile_id FROM claim_revisions cr
                    JOIN claims c ON c.id=cr.claim_id WHERE cr.id=?""", (revision_id,)).fetchone()
                if not row or row["profile_id"] != track["profile_id"]:
                    raise ValueError("Supporting claim belongs to another profile or does not exist")
            columns = ",".join(TRACK_FIELDS)
            placeholders = ",".join("?" for _ in TRACK_FIELDS)
            version = 1 if not previous else previous["version_number"] + 1
            return db.execute(f"""INSERT INTO research_track_versions
                (track_id,version_number,{columns},supporting_claim_revision_ids_json,created_at)
                VALUES(?,?,{placeholders},?,?)""", (
                track_id, version, *(values[field] for field in TRACK_FIELDS),
                _json(list(dict.fromkeys(claim_ids))), utc_now(),
            )).lastrowid

    def approve_track(self, version_id: int, reviewer: str) -> None:
        if not reviewer.strip():
            raise ValueError("Reviewer is required")
        with transaction(self.db_path) as db:
            version = db.execute("SELECT * FROM research_track_versions WHERE id=?", (version_id,)).fetchone()
            if not version or version["approval_state"] != "DRAFT":
                raise ValueError("Only draft track versions can be approved")
            for revision_id in json.loads(version["supporting_claim_revision_ids_json"]):
                claim = db.execute("SELECT review_status FROM claim_revisions WHERE id=?", (revision_id,)).fetchone()
                if not claim or claim["review_status"] != "APPROVED":
                    raise ValueError("Supporting applicant claims must be approved first")
            db.execute("""UPDATE research_track_versions SET approval_state='APPROVED',
                approved_at=?,approved_by=? WHERE id=?""", (utc_now(), reviewer.strip(), version_id))

    def return_track_to_draft(self, track_id: int) -> int:
        """Preserve the reviewed version and open a new editable draft."""
        return self.revise_track(track_id)

    def update_track(self, track_id: int, *, title: str | None = None,
                     priority: int | None = None, archive: bool | None = None) -> None:
        with transaction(self.db_path) as db:
            track = db.execute("SELECT * FROM research_tracks WHERE id=?", (track_id,)).fetchone()
            if not track:
                raise ValueError("Track does not exist")
            if priority is not None and not 1 <= priority <= 5:
                raise ValueError("Priority must be 1–5")
            db.execute("""UPDATE research_tracks SET title=?,priority=?,status=?,updated_at=? WHERE id=?""", (
                title.strip() if title is not None else track["title"],
                priority if priority is not None else track["priority"],
                ("ARCHIVED" if archive else "ACTIVE") if archive is not None else track["status"],
                utc_now(), track_id,
            ))

    def list_tracks(self, profile_id: int) -> list[dict]:
        with connect(self.db_path) as db:
            return [dict(r) for r in db.execute("""SELECT t.*, v.id AS version_id,
                v.version_number,v.approval_state,v.research_problem,v.proposed_methodology,
                v.supporting_claim_revision_ids_json
                FROM research_tracks t JOIN research_track_versions v ON v.track_id=t.id
                WHERE t.profile_id=? AND v.version_number=(SELECT MAX(x.version_number)
                    FROM research_track_versions x WHERE x.track_id=t.id)
                ORDER BY t.status,t.priority,t.id""", (profile_id,))]

    def track_history(self, track_id: int) -> list[dict]:
        with connect(self.db_path) as db:
            return [dict(r) for r in db.execute("""SELECT * FROM research_track_versions
                WHERE track_id=? ORDER BY version_number DESC""", (track_id,))]
