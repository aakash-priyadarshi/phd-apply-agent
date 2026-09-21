import io
import sqlite3
from types import SimpleNamespace

import pytest
from docx import Document

from phd_agent.db import connect
from phd_agent.documents import DocumentVault
from phd_agent.profile import ApplicantTruth, CandidateBatch


@pytest.fixture
def truth(tmp_path):
    return ApplicantTruth(tmp_path / "app.db")


def test_claim_review_and_immutable_profile_snapshot(truth):
    profile = truth.create_profile("Applicant")
    aspiration = truth.create_claim(profile, "OTHER", "I aim to study reliable AI",
                                     "future_direction", "ASPIRATION")
    truth.review_claim(aspiration, True, "Applicant", for_application=True)
    with pytest.raises(ValueError, match="linked evidence"):
        fact = truth.create_claim(profile, "SKILL_METHOD", "I used Python", "method", "FACT")
        truth.review_claim(fact, True, "Applicant", verification_state="VERIFIED")
    version = truth.create_profile_version(profile, [aspiration], approve=True, reviewer="Applicant")
    snapshot = truth.profile_snapshot(version)
    assert snapshot[0]["claim_text"] == "I aim to study reliable AI"
    with connect(truth.db_path) as db, pytest.raises(sqlite3.IntegrityError):
        db.execute("UPDATE profile_versions SET notes='changed' WHERE id=?", (version,))
    with connect(truth.db_path) as db, pytest.raises(sqlite3.IntegrityError):
        db.execute("DELETE FROM profile_versions WHERE id=?", (version,))
    with connect(truth.db_path) as db, pytest.raises(sqlite3.IntegrityError):
        db.execute("UPDATE profile_version_claims SET claim_snapshot_json='{}' WHERE profile_version_id=?", (version,))
    revision = truth.revise_claim(snapshot[0]["claim_id"], claim_text="I plan to study reliable AI")
    assert truth.profile_snapshot(version) == snapshot
    assert truth.list_claims(profile)[0]["id"] == revision
    assert truth.list_claims(profile)[0]["review_status"] == "PENDING"
    with connect(truth.db_path) as db, pytest.raises(sqlite3.IntegrityError):
        db.execute("DELETE FROM claim_revisions WHERE id=?", (aspiration,))


def test_fact_inference_and_aspiration_rules(truth):
    profile = truth.create_profile("Applicant")
    vault = DocumentVault(truth.db_path)
    evidence = vault.upload(b"source", "evidence.txt", "SOURCE", "CERTIFICATE", "Evidence")
    fact = truth.create_claim(profile, "EDUCATION", "Completed degree", "degree", "FACT",
                              source_document_version_id=evidence["version_id"])
    with pytest.raises(ValueError, match="source document version must be approved"):
        truth.review_claim(fact, True, "Applicant", verification_state="VERIFIED")
    vault.set_approval(evidence["version_id"], True, "Applicant")
    with pytest.raises(ValueError, match="verified"):
        truth.review_claim(fact, True, "Applicant")
    truth.review_claim(fact, True, "Applicant", verification_state="VERIFIED", for_application=True)
    inference = truth.create_claim(profile, "OTHER", "This suggests strong systems experience",
                                   "inference", "INFERENCE", source_document_version_id=evidence["version_id"])
    truth.review_claim(inference, True, "Applicant", for_outreach=True)
    assert truth.list_claims(profile)[0]["approved_for_application"] == 1
    assert truth.list_claims(profile)[1]["approved_for_outreach"] == 1
    false_aspiration = truth.create_claim(profile, "OTHER", "I developed novel RL agents", "goal", "ASPIRATION")
    with pytest.raises(ValueError, match="future aim"):
        truth.review_claim(false_aspiration, True, "Applicant")
    assert truth._aspiration_wording("I seek to investigate reliable agent evaluation.")
    assert truth._aspiration_wording("My objective is to investigate reliable agent evaluation.")
    assert not truth._aspiration_wording("My objective was to investigate reliable agent evaluation.")
    assert not truth._aspiration_wording("I developed AI planning systems.")


def test_docx_ingestion_and_structured_candidates_reject_unsupported(truth):
    profile = truth.create_profile("Applicant")
    doc = Document()
    doc.add_paragraph("Research project: evaluated a retrieval model.")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Result: 20 tests completed."
    data = io.BytesIO()
    doc.save(data)
    vault = DocumentVault(truth.db_path)
    uploaded = vault.upload(data.getvalue(), "cv.docx", "SOURCE", "CV", "Current CV")
    with pytest.raises(ValueError, match="Approve"):
        truth.ingest_document(uploaded["version_id"])
    vault.set_approval(uploaded["version_id"], True, "Applicant")
    extraction_id, pages = truth.ingest_document(uploaded["version_id"])
    assert len(pages) == 1 and "20 tests" in pages[0]
    good = {"statement": "Evaluated a retrieval model", "category": "RESEARCH_PROJECT",
            "normalized_claim_type": "project", "classification": "FACT",
            "source_quote": "evaluated a retrieval model", "source_location": "paragraph 1",
            "confidence": 0.8, "fields": []}
    bad = {**good, "statement": "Published in Nature", "source_quote": "published in Nature"}
    batch = CandidateBatch.model_validate({"candidates": [good, bad]})
    client = SimpleNamespace(responses=SimpleNamespace(parse=lambda **_: SimpleNamespace(output_parsed=batch)))
    result = truth.extract_candidates(profile, extraction_id, api_key="", client=client)
    assert result == {"pending": 1, "rejected_unsupported": 1}
    claims = truth.list_claims(profile)
    assert [c["review_status"] for c in claims] == ["PENDING", "REJECTED"]
    assert all(c["approved_for_application"] == 0 for c in claims)


def test_pdf_ingestion_preserves_every_page(truth, monkeypatch):
    vault = DocumentVault(truth.db_path)
    uploaded = vault.upload(b"pdf fixture", "cv.pdf", "SOURCE", "CV", "PDF CV")
    vault.set_approval(uploaded["version_id"], True, "Applicant")
    class Page:
        def __init__(self, text):
            self.text = text
        def extract_text(self):
            return self.text
    monkeypatch.setattr("phd_agent.profile.PdfReader", lambda _: SimpleNamespace(
        pages=[Page("Education section"), Page("Research project on page two")]))
    extraction_id, pages = truth.ingest_document(uploaded["version_id"])
    assert pages == ["Education section", "Research project on page two"]
    with connect(truth.db_path) as db:
        record = db.execute("SELECT * FROM applicant_document_extractions WHERE id=?", (extraction_id,)).fetchone()
    assert record["page_count"] == 2
    assert "Research project on page two" in record["extracted_text"]


def test_track_versioning_and_approved_immutability(truth):
    profile = truth.create_profile("Applicant")
    track_id, v1 = truth.create_track(profile, "Reliable agents", research_problem="How to evaluate agents?")
    truth.approve_track(v1, "Applicant")
    with connect(truth.db_path) as db, pytest.raises(sqlite3.IntegrityError):
        db.execute("UPDATE research_track_versions SET research_problem='changed' WHERE id=?", (v1,))
    with connect(truth.db_path) as db, pytest.raises(sqlite3.IntegrityError):
        db.execute("DELETE FROM research_track_versions WHERE id=?", (v1,))
    v2 = truth.revise_track(track_id, research_problem="How to evaluate grounded agents?")
    assert [r["version_number"] for r in truth.track_history(track_id)] == [2, 1]
    assert truth.track_history(track_id)[0]["approval_state"] == "DRAFT"
    truth.update_track(track_id, priority=1, archive=True)
    assert truth.list_tracks(profile)[0]["status"] == "ARCHIVED"


def test_overlapping_chunks_cover_boundaries_and_short_pages():
    from phd_agent.profile import _overlapping_chunks
    assert _overlapping_chunks("") == []
    assert _overlapping_chunks("short") == ["short"]
    text = "a" * 15000
    chunks = _overlapping_chunks(text, size=10000, overlap=1000)
    assert chunks[0] == text[:10000]
    assert chunks[1] == text[9000:15000]
    assert chunks[0][-1000:] == chunks[1][:1000]
